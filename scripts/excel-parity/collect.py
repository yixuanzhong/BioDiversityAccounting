#!/usr/bin/env python3
"""Read RESULTS and enhancement rows from LibreOffice-recalculated workbooks.

    python3 scripts/excel-parity/collect.py WORKDIR     # reads WORKDIR/out/s<i>.xlsx
"""
import json
import sys
from pathlib import Path

import openpyxl

work = Path(sys.argv[1])
n = len(json.loads((work / "scenarios.json").read_text()))
cells = [("tb", "F12"), ("fb", "F13"), ("mb", "F14"), ("tp", "F15"), ("fp", "F16"), ("mp", "F17"), ("te", "E30"), ("fe", "E31"), ("me", "E32")]
enh_sheets = [("terrestrial", "A3 Enhancement - Terrestrial"), ("freshwater", "B3 Enhancement - FW"), ("marine", "C3 Enhancement - Marine")]
res = []
for i in range(n):
    wb = openpyxl.load_workbook(work / "out" / f"s{i}.xlsx", data_only=True)
    d = {k: wb["RESULTS"][c].value for k, c in cells}
    # (years, multiplier, units) per enhancement row, aligned with baseline rows
    d["enh"] = {k: [(wb[s].cell(r, 20).value, wb[s].cell(r, 21).value, wb[s].cell(r, 23).value) for r in range(7, 16)] for k, s in enh_sheets}
    res.append(d)
(work / "excel_results.json").write_text(json.dumps(res, default=str))
print(f"collected {n} results")
