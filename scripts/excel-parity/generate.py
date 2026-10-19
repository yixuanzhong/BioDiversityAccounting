#!/usr/bin/env python3
"""Write random assessments into copies of the official Metric 1.1 workbook.

    python3 scripts/excel-parity/generate.py METRIC.xlsx WORKDIR [N]

Creates WORKDIR/scenarios.json and WORKDIR/in/s<i>.xlsx. Recalculate the copies with
LibreOffice, then run collect.py and compare.js (see README).
"""
import json
import random
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SHEETS = {
    "terrestrial": ("A1 Baseline - Terrestrial", "A2 Creation - Terrestrial", "A3 Enhancement - Terrestrial"),
    "freshwater": ("B1 Baseline - FW", "B2 Creation - FW", "B3 Enhancement - FW"),
    "marine": ("C1 Baseline - Marine", "C2 Creation - Marine", "C3 Enhancement - Marine"),
}
CONDS = ["Good", "Moderate", "Poor"]


def load_data():
    js = (ROOT / "assets/js/data.js").read_text()
    return json.loads(js[js.index("var DATA = ") + 11: js.rindex("};") + 1])


def scenario(seed, data):
    rnd = random.Random(seed)
    out = {}
    for eco in SHEETS:
        habs = [h for h in data["habitats"] if h["eco"] == eco]
        base, crea, enh = [], [], {}
        for i in range(rnd.randint(2, 8)):
            h = rnd.choice(habs)
            area = round(rnd.uniform(0.1, 20), 4)
            cond = rnd.choice(CONDS)
            ret = round(area * rnd.choice([0, 0, 0.3, 0.5]), 4)
            en_area = round((area - ret) * rnd.choice([0, 0.4, 1]), 4) if rnd.random() < 0.7 else 0
            rid = f"{eco}{i}"
            base.append(dict(id=rid, habitat=h["id"], area=area, condition=cond, retained=ret, enhanced=en_area))
            if en_area > 0:
                opts = data["enhance"].get((h["name"] + "/" + cond).upper(), {})
                if opts and rnd.random() < 0.8:
                    tname, tcond = rnd.choice(list(opts)).rsplit("/", 1)
                    t = next(x for x in habs if x["name"].upper() == tname)
                    enh[rid] = dict(target=t["id"], condition=tcond.capitalize())
                else:
                    enh[rid] = dict(target=rnd.choice(habs)["id"], condition=rnd.choice(CONDS))
        for i in range(rnd.randint(1, 6)):
            h = rnd.choice(habs)
            crea.append(dict(id=f"{eco}c{i}", habitat=h["id"], area=round(rnd.uniform(0.1, 15), 4), condition=rnd.choice(CONDS)))
        out[eco] = dict(baseline=base, creation=crea, enhancement=enh)
    return out


def write(metric, sc, path):
    wb = openpyxl.load_workbook(metric)
    name = lambda hid: hid.split(":", 1)[1]
    for eco, (b, c, e) in SHEETS.items():
        for i, r in enumerate(sc[eco]["baseline"]):
            ws, row = wb[b], 8 + i
            for col, v in ((5, name(r["habitat"])), (6, r["area"]), (9, r["condition"]), (14, r["retained"]), (15, r["enhanced"])):
                ws.cell(row, col).value = v
            en = sc[eco]["enhancement"].get(r["id"])
            if en:
                wb[e].cell(7 + i, 8).value = name(en["target"])
                wb[e].cell(7 + i, 15).value = en["condition"]
        for i, r in enumerate(sc[eco]["creation"]):
            for col, v in ((5, name(r["habitat"])), (6, r["area"]), (9, r["condition"])):
                wb[c].cell(7 + i, col).value = v
    # LibreOffice < 24.8 lacks XMATCH; MATCH(x, range, 0) is equivalent for these exact-match lookups.
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                t = cell.value.text if hasattr(cell.value, "text") else cell.value
                if isinstance(t, str) and "XMATCH" in t:
                    cell.value = re.sub(r"_xlfn\.XMATCH\(([^,]+),([^)]+)\)", r"MATCH(\1,\2,0)", t)
    wb.save(path)


def main():
    metric, work = sys.argv[1], Path(sys.argv[2])
    n = int(sys.argv[3]) if len(sys.argv) > 3 else 20
    data = load_data()
    (work / "in").mkdir(parents=True, exist_ok=True)
    scen = [scenario(s, data) for s in range(n)]
    (work / "scenarios.json").write_text(json.dumps(scen))
    for i, s in enumerate(scen):
        write(metric, s, work / "in" / f"s{i}.xlsx")
    print(f"wrote {n} workbooks to {work / 'in'}")


if __name__ == "__main__":
    main()
